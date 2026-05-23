import json
import logging
import os
import threading

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from fastapi import FastAPI, Request, Form, BackgroundTasks, HTTPException
from fastapi.responses import HTMLResponse, JSONResponse, FileResponse
from fastapi.templating import Jinja2Templates
import uvicorn

from db import (init_db, reset_stale_campaigns, create_campaign, get_campaign,
                get_all_campaigns, get_leads, get_lead, get_setting, set_setting,
                update_campaign, count_running_campaigns)
from runner import run_campaign, send_campaign_emails, stop_campaign, analyze_campaign_leads

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
logger = logging.getLogger(__name__)

app = FastAPI()
templates = Jinja2Templates(directory="templates")
templates.env.filters["fromjson"] = lambda s: json.loads(s) if s else []

SECTORS = [
    "restaurant", "cafe", "hotel", "friseur", "zahnarzt",
    "rechtsanwalt", "steuerberater", "fitnessstudio", "apotheke",
    "autohaus", "immobilien", "architekt", "fotograf", "marketing",
    "physiotherapie", "kosmetik", "blumenladen", "buchhandlung",
    "elektronik", "fahrschule", "reinigung", "umzug", "catering",
    "baecker", "metzger", "optiker", "versicherung", "tierarzt",
]

CITIES = {
    "berlin":      "Berlin",
    "hamburg":     "Hamburg",
    "muenchen":    "München",
    "frankfurt":   "Frankfurt",
    "koeln":       "Köln",
    "duesseldorf": "Düsseldorf",
    "stuttgart":   "Stuttgart",
    "leipzig":     "Leipzig",
    "dresden":     "Dresden",
    "nuernberg":   "Nürnberg",
}

SECTORS_SET = set(SECTORS)
CITIES_SET  = set(CITIES)
MAX_CONCURRENT_CAMPAIGNS = 3


@app.on_event("startup")
async def startup():
    os.makedirs("output", exist_ok=True)
    init_db()
    reset_stale_campaigns()
    logger.info("Uygulama başlatıldı")


def _get_all_sectors() -> list[str]:
    """Sabit sektörler + kullanıcının özel sektörleri."""
    custom_raw = get_setting("custom_sectors", "")
    custom = [s.strip() for s in custom_raw.splitlines() if s.strip()]
    seen = set(SECTORS)
    extra = [s for s in custom if s not in seen]
    return SECTORS + extra


@app.get("/", response_class=HTMLResponse)
async def dashboard(request: Request):
    campaigns = get_all_campaigns()
    return templates.TemplateResponse(
        request=request, name="index.html",
        context={
            "campaigns": [dict(c) for c in campaigns],
            "sectors":   _get_all_sectors(),
            "cities":    CITIES,
        }
    )


@app.post("/campaign/start")
async def start_campaign(
    city:        str       = Form(...),
    sectors:     list[str] = Form(...),
    home_office: bool      = Form(False),
    analyze:     bool      = Form(False),
):
    if city not in CITIES_SET:
        raise HTTPException(status_code=422, detail=f"Geçersiz şehir: {city}")

    if not sectors:
        raise HTTPException(status_code=422, detail="En az bir sektör seçilmeli")

    all_valid = set(_get_all_sectors())
    invalid = [s for s in sectors if s not in all_valid]
    if invalid:
        raise HTTPException(status_code=422, detail=f"Geçersiz sektörler: {invalid}")

    if count_running_campaigns() >= MAX_CONCURRENT_CAMPAIGNS:
        raise HTTPException(
            status_code=429,
            detail=f"Zaten {MAX_CONCURRENT_CAMPAIGNS} kampanya çalışıyor. Lütfen bekle."
        )

    search_mode  = "home_office" if home_office else "standard"
    analyze_mode = "find_analyze" if analyze else "find_only"
    cid = create_campaign(city, sectors, search_mode=search_mode, analyze_mode=analyze_mode)
    threading.Thread(
        target=run_campaign,
        args=(cid, city, sectors, home_office, analyze),
        daemon=True,
        name=f"campaign-{cid}",
    ).start()
    logger.info("Kampanya %d başlatıldı — %s, ev ofisi=%s, analiz=%s", cid, city, home_office, analyze)
    return JSONResponse({"campaign_id": cid})


@app.get("/campaign/{cid}", response_class=HTMLResponse)
async def campaign_page(request: Request, cid: int):
    campaign = get_campaign(cid)
    if not campaign:
        raise HTTPException(status_code=404, detail="Kampanya bulunamadı")
    leads = get_leads(cid)
    can_analyze = (
        campaign["status"] in ("completed", "stopped") and
        any(l["website"] and not l["analyzed"] for l in leads)
    )
    return templates.TemplateResponse(
        request=request, name="campaign.html",
        context={
            "campaign":     dict(campaign),
            "leads":        [dict(l) for l in leads],
            "can_analyze":  can_analyze,
        }
    )


@app.get("/campaign/{cid}/progress")
async def campaign_progress(cid: int):
    c = get_campaign(cid)
    if not c:
        return JSONResponse({"status": "not_found"})
    return JSONResponse({
        "status": c["status"],
        "step":   c["progress_step"],
        "pct":    c["progress_pct"],
        "total":  c["total_found"],
    })


@app.post("/campaign/{cid}/stop")
async def stop_campaign_endpoint(cid: int):
    campaign = get_campaign(cid)
    if not campaign:
        raise HTTPException(status_code=404, detail="Kampanya bulunamadı")
    if campaign["status"] not in ("running", "analyzing"):
        raise HTTPException(status_code=400, detail="Kampanya zaten çalışmıyor")
    ok = stop_campaign(cid)
    if not ok:
        update_campaign(cid, status="stopped", progress_step="Durduruldu")
    return JSONResponse({"ok": True})


@app.post("/campaign/{cid}/analyze")
async def start_analysis(cid: int):
    campaign = get_campaign(cid)
    if not campaign:
        raise HTTPException(status_code=404, detail="Kampanya bulunamadı")
    if campaign["status"] not in ("completed", "stopped"):
        raise HTTPException(status_code=400, detail="Kampanya henüz tamamlanmadı")
    threading.Thread(
        target=analyze_campaign_leads,
        args=(cid,),
        daemon=True,
        name=f"analyze-{cid}",
    ).start()
    logger.info("Manuel analiz başlatıldı — kampanya %d", cid)
    return JSONResponse({"ok": True})


@app.post("/campaign/{cid}/send-emails")
async def send_emails(cid: int, background_tasks: BackgroundTasks):
    campaign = get_campaign(cid)
    if not campaign:
        raise HTTPException(status_code=404, detail="Kampanya bulunamadı")
    if campaign["status"] not in ("completed", "stopped"):
        raise HTTPException(status_code=400, detail="Kampanya henüz tamamlanmadı")
    background_tasks.add_task(send_campaign_emails, cid)
    return JSONResponse({"ok": True})


@app.get("/campaign/{cid}/download")
async def download_excel(cid: int):
    campaign = get_campaign(cid)
    if not campaign:
        raise HTTPException(status_code=404, detail="Kampanya bulunamadı")

    leads = get_leads(cid)
    path  = f"output/campaign_{cid}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leads"

    headers = ["Öncelik", "İşletme Adı", "Sektör", "Website", "Email",
               "Puan", "Sorunlar", "Mail Gönderildi"]
    hfill = PatternFill(start_color="1A252F", end_color="1A252F", fill_type="solid")
    hfont = Font(color="FFFFFF", bold=True)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = hfill
        cell.font = hfont
        cell.alignment = Alignment(horizontal="center")

    fills = {
        "YUKSEK": PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid"),
        "ORTA":   PatternFill(start_color="FDEBD0", end_color="FDEBD0", fill_type="solid"),
        "DUSUK":  PatternFill(start_color="D5F5E3", end_color="D5F5E3", fill_type="solid"),
    }

    for r, lead in enumerate(leads, 2):
        row = [lead["priority"], lead["name"], lead["sector"], lead["website"],
               lead["email"], lead["quality_score"], lead["issues"], lead["email_sent"]]
        f = fills.get(lead["priority"] or "ORTA", fills["ORTA"])
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.fill = f
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    widths = [10, 30, 15, 40, 30, 8, 50, 14]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    wb.save(path)
    return FileResponse(path, filename=f"leads_{campaign['city']}_{cid}.xlsx")


@app.get("/campaign/{cid}/lead/{lid}/preview")
async def lead_preview(cid: int, lid: int):
    campaign = get_campaign(cid)
    if not campaign:
        raise HTTPException(status_code=404, detail="Kampanya bulunamadı")
    lead = get_lead(lid)
    if not lead or lead["campaign_id"] != cid:
        raise HTTPException(status_code=404, detail="Lead bulunamadı")

    lead_dict = dict(lead)
    template_de = get_setting("email_template_de", "")
    template_en = get_setting("email_template_en", "")
    has_key = bool(get_setting("anthropic_api_key", ""))

    email_preview = None
    if template_de or template_en:
        issues_text = lead_dict.get("issues") or "—"
        vars_ = dict(
            name=lead_dict.get("name", ""),
            sector=lead_dict.get("sector", ""),
            website=lead_dict.get("website") or "—",
            issues=issues_text,
            score=lead_dict.get("quality_score") or "?",
        )
        try:
            de = template_de.format_map(vars_) if template_de else ""
            en = template_en.format_map(vars_) if template_en else ""
            parts = []
            if de: parts.append(f"🇩🇪 Almanca:\n{de}")
            if en: parts.append(f"🇬🇧 İngilizce:\n{en}")
            email_preview = "\n\n---\n\n".join(parts)
        except (KeyError, ValueError):
            email_preview = template_de or template_en

    return JSONResponse({
        "lead": lead_dict,
        "email_preview": email_preview,
        "has_template": bool(template_de or template_en),
        "has_api_key": has_key,
    })


@app.get("/settings", response_class=HTMLResponse)
async def settings_page(request: Request):
    return templates.TemplateResponse(
        request=request, name="settings.html",
        context={
            "anthropic_key":     get_setting("anthropic_api_key"),
            "gmail_user":        get_setting("gmail_user"),
            "gmail_pass":        get_setting("gmail_password"),
            "email_template_de": get_setting("email_template_de"),
            "email_template_en": get_setting("email_template_en"),
            "custom_sectors":    get_setting("custom_sectors"),
        }
    )


@app.post("/settings")
async def save_settings(
    anthropic_key:    str = Form(""),
    gmail_user:       str = Form(""),
    gmail_pass:       str = Form(""),
    email_template_de: str = Form(""),
    email_template_en: str = Form(""),
    custom_sectors:    str = Form(""),
):
    set_setting("anthropic_api_key",  anthropic_key.strip())
    set_setting("gmail_user",         gmail_user.strip())
    set_setting("gmail_password",     gmail_pass.strip())
    set_setting("email_template_de",  email_template_de.strip())
    set_setting("email_template_en",  email_template_en.strip())
    set_setting("custom_sectors",     custom_sectors.strip())
    logger.info("Ayarlar güncellendi")
    return JSONResponse({"ok": True})


if __name__ == "__main__":
    uvicorn.run("app:app", host="127.0.0.1", port=8000, reload=True)
