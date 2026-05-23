import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import os


def save_to_excel(businesses: list[dict], filepath: str) -> str:
    os.makedirs(os.path.dirname(filepath), exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Berlin Leads"

    headers = [
        "Öncelik", "İşletme Adı", "Sektör", "Website", "Email",
        "Telefon", "Kalite Puanı", "Sorunlar", "Yükleme Süresi",
        "Kaynak", "Mail Gönderildi", "Notlar"
    ]

    header_fill = PatternFill(start_color="1A252F", end_color="1A252F", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    sorted_biz = sorted(businesses, key=lambda x: x.get("quality_score", 100))

    red = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")
    yellow = PatternFill(start_color="FDEBD0", end_color="FDEBD0", fill_type="solid")
    green = PatternFill(start_color="D5F5E3", end_color="D5F5E3", fill_type="solid")

    for row_i, biz in enumerate(sorted_biz, 2):
        score = biz.get("quality_score", 100)
        if score < 40:
            priority = "YUKSEK"
            fill = red
        elif score < 70:
            priority = "ORTA"
            fill = yellow
        else:
            priority = "DUSUK"
            fill = green

        row = [
            priority,
            biz.get("name", ""),
            biz.get("sector", ""),
            biz.get("website", ""),
            biz.get("email", ""),
            biz.get("phone", ""),
            score,
            ", ".join(biz.get("issues", [])),
            biz.get("load_time", ""),
            biz.get("source", ""),
            biz.get("email_sent", "Hayır"),
            "",
        ]

        for col, val in enumerate(row, 1):
            cell = ws.cell(row=row_i, column=col, value=val)
            cell.fill = fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    col_widths = [10, 30, 15, 40, 32, 16, 14, 55, 14, 14, 16, 20]
    for col, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.freeze_panes = "A2"
    wb.save(filepath)
    print(f"Excel kaydedildi: {filepath} ({len(sorted_biz)} kayıt)")
    return filepath
